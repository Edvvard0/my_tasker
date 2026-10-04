"""Statement parsing: CSV, XLSX and PDF samples (synthetic), header detection, row rules."""

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import pytest
from fpdf import FPDF
from hypothesis import given
from hypothesis import strategies as st

from tasker.banks import reference as ref
from tasker.banks import statements as st_mod
from tasker.banks.statements import StatementError, parse_statement

DATA = Path(__file__).resolve().parent / "data" / "banks"
CSV = (DATA / "tbank_synthetic.csv").read_bytes()
PDF = (DATA / "vtb_synthetic.pdf").read_bytes()


def candidates(result: dict[str, Any]) -> list[dict[str, Any]]:
    found: list[dict[str, Any]] = result["candidates"]
    return found


# ------------------------------------------------------------------ the synthetic T-Bank CSV


def test_tbank_csv_candidates() -> None:
    result = parse_statement(CSV)
    assert (result["format"], result["bank"], result["cards"]) == ("csv", "tbank", ["1234"])
    assert result["period"] == {"from": "2026-10-03", "to": "2026-10-07"}
    assert result["skipped"] == [{"row": 5, "reason": "status failed"}]  # the FAILED payment
    rows = candidates(result)
    assert [(c["kind"], c["amount"]) for c in rows] == [
        ("expense", 123_456),
        ("expense", 34_900),
        ("expense", 34_900),
        ("income", 8_500_000),
        ("expense", 115_000),
        ("income", 30_000),
    ]
    first = rows[0]
    assert first["occurred_at"] == "2026-10-03T08:30:12Z"  # Moscow time -> UTC
    assert (first["date_only"], first["merchant_norm"], first["mcc"]) == (
        False,
        "пятерочка",
        "5411",
    )
    assert first["bank_category"] == "Супермаркеты"
    assert first["suggested_category"]["system_key"] == "expense.groceries"


def test_two_identical_operations_get_different_tails() -> None:
    first, second = candidates(parse_statement(CSV))[1:3]
    assert first["dedup_tail"] == "expense|34900|2026-10-03T11:05|yandex taxi"
    assert second["dedup_tail"] == first["dedup_tail"] + "|1"
    assert ref.hash_of("a", first["dedup_tail"]) != ref.hash_of("a", second["dedup_tail"])


def test_foreign_currency_keeps_the_account_amount_and_asks_for_review() -> None:
    usd = candidates(parse_statement(CSV))[4]
    assert (usd["amount"], usd["currency"]) == (115_000, "RUB")
    assert (usd["original_amount"], usd["original_currency"]) == (1_250, "USD")
    assert (usd["needs_review"], usd["review_reason"]) == (True, "foreign_currency")


def test_cp1251_and_utf8_with_bom_and_other_delimiters() -> None:
    text = CSV.decode("utf-8")
    cp1251 = parse_statement(text.encode("cp1251"))
    assert cp1251["candidates"] == parse_statement(CSV)["candidates"]
    bom = parse_statement(b"\xef\xbb\xbf" + CSV)
    assert bom["candidates"] == parse_statement(CSV)["candidates"]
    tabs = parse_statement(text.replace(";", "\t").encode())
    assert len(tabs["candidates"]) == 6
    commas = "\n".join(
        ",".join(f'"{cell}"' for cell in line.split(";")) for line in text.splitlines()
    )
    assert len(parse_statement(commas.encode())["candidates"]) == 6


def test_the_bank_can_be_forced_and_unknown_headers_fall_back_to_generic() -> None:
    assert parse_statement(CSV, bank="generic")["bank"] in {"generic", "tbank"}
    assert parse_statement(CSV, bank="vtb")["bank"] in {"generic", "vtb"}
    rows = b"Date;Amount;Description\n2026-10-03 10:00:00;-10.5;Shop\n"
    result = parse_statement(rows)
    assert result["bank"] == "generic"
    assert candidates(result)[0]["amount"] == 1_050


# ------------------------------------------------------------------ the synthetic VTB PDF


def test_pdf_table_and_the_lines_around_it() -> None:
    result = parse_statement(PDF)
    assert (result["format"], result["bank"]) == ("pdf", "vtb")
    assert result["period"] == {"from": "2026-10-01", "to": "2026-10-07"}
    assert result["closing_balance"] == {"amount": 10_201_545, "at": "2026-10-07T20:59:59Z"}
    rows = candidates(result)
    assert [(c["kind"], c["amount"], c["date_only"]) for c in rows] == [
        ("expense", 123_456, True),
        ("income", 500_000, True),
        ("expense", 95_000, True),
        ("expense", 79_999, True),
    ]
    assert rows[0]["occurred_at"] == "2026-10-03T09:00:00Z"  # a date only: 12:00 Moscow
    assert rows[0]["balance_after"] == 9_876_544
    steam = rows[2]
    assert (steam["original_amount"], steam["original_currency"], steam["needs_review"]) == (
        1_000,
        "USD",
        True,
    )


def test_pdf_without_a_table_is_unrecognised() -> None:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 8, "Just a letter, no table")
    with pytest.raises(StatementError) as raised:
        parse_statement(bytes(pdf.output()))
    assert raised.value.code == "statement_unrecognized"


def test_pdf_garbage_is_unreadable() -> None:
    with pytest.raises(StatementError) as raised:
        parse_statement(b"%PDF-1.4 this is not a pdf at all")
    assert raised.value.code == "statement_unreadable"


# ------------------------------------------------------------------ XLSX


def _workbook(rows: list[list[Any]], extra_sheet: bool = False) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    if extra_sheet:
        sheet.title = "Cover"
        sheet.append(["Nothing useful here"])
        sheet = book.create_sheet("Operations")
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


VTB_HEADER = [
    "Дата операции",
    "Описание операции",
    "Расход",
    "Приход",
    "Валюта операции",
    "Остаток",
    "Номер операции",
]


def test_xlsx_with_real_dates_numbers_and_debit_credit_columns() -> None:
    data = _workbook(
        [
            ["Выписка ВТБ (синтетический образец)"],
            ["за период с 01.10.2026 по 31.10.2026"],
            ["Остаток на конец периода: 1 234,50 RUB"],
            VTB_HEADER,
            [datetime(2026, 10, 3, 12, 15, 0), "Оплата Магнит", 349.9, None, "RUB", 100.5, "OP-1"],
            [datetime(2026, 10, 4), "Зарплата", None, 85000, "RUB", 85100.5, "OP-2"],
            [datetime(2026, 10, 5), "Оплата Steam", 12, None, "USD", None, "OP-3"],
            [datetime(2026, 10, 6), "Нулевая", 0, 0, "RUB", None, "OP-4"],
            [None, None, None, None, None, None, None],
            ["Итого", None, 361.9, 85000, None, None, None],
        ]
    )
    result = parse_statement(data)
    assert (result["format"], result["bank"]) == ("xlsx", "vtb")
    assert result["period"] == {"from": "2026-10-01", "to": "2026-10-31"}
    assert result["closing_balance"] == {"amount": 123_450, "at": "2026-10-31T20:59:59Z"}
    rows = candidates(result)
    assert [(c["kind"], c["amount"], c["external_id"]) for c in rows] == [
        ("expense", 34_990, "OP-1"),
        ("income", 8_500_000, "OP-2"),
        ("expense", 1_200, "OP-3"),
    ]
    assert rows[0]["occurred_at"] == "2026-10-03T09:15:00Z"
    assert rows[0]["date_only"] is False
    assert rows[1]["date_only"] is True  # midnight in a workbook means "no time"
    assert (rows[2]["needs_review"], rows[2]["currency"]) == (True, "USD")
    assert {"row": 8, "reason": "no_amount"} in result["skipped"]
    assert rows[0]["balance_after"] == 10_050


def test_xlsx_looks_through_the_sheets_for_the_table() -> None:
    data = _workbook([VTB_HEADER, ["04.10.2026", "Кофе", "100,00", "", "RUB", "", ""]], True)
    assert [c["amount"] for c in candidates(parse_statement(data))] == [10_000]


def test_xlsx_without_a_table_and_broken_xlsx() -> None:
    with pytest.raises(StatementError) as raised:
        parse_statement(_workbook([["a", "b"], ["c", "d"]]))
    assert raised.value.code == "statement_unrecognized"
    with pytest.raises(StatementError) as broken:
        parse_statement(b"PK\x03\x04 not really a zip")
    assert broken.value.code == "statement_unreadable"


def test_xlsx_too_large_unpacked(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(st_mod, "MAX_UNPACKED_BYTES", 10)
    with pytest.raises(StatementError) as raised:
        parse_statement(_workbook([VTB_HEADER]))
    assert raised.value.code == "statement_too_large"


def _forged_dimension(data: bytes, ref: str) -> bytes:
    """The same workbook with the ``<dimension>`` of its first sheet replaced by ``ref``."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
        for info in src.infolist():
            content = src.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content, count = re.subn(
                    rb'<dimension ref="[^"]*"', f'<dimension ref="{ref}"'.encode(), content
                )
                assert count == 1
            dst.writestr(info, content)
    return out.getvalue()


def test_xlsx_with_a_forged_dimension_is_not_padded() -> None:
    rows = [VTB_HEADER] + [
        ["03.10.2026", f"Магазин {n}", "10,00", "", "RUB", "", ""] for n in range(300)
    ]
    forged = _forged_dimension(_workbook(rows), "A1:XFD1048576")
    (table,) = st_mod.extract_tables("xlsx", forged)
    assert len(table) == 301
    assert max(len(row) for row in table) == len(VTB_HEADER)  # no padding to 16 384 columns
    assert len(candidates(parse_statement(forged))) == 300


def test_xlsx_cells_beyond_the_column_limit_are_cut() -> None:
    wide = [*VTB_HEADER, *["x"] * 70, "far away"]
    row = ["03.10.2026", "Кофе", "100,00", "", "RUB", "", ""]
    (table,) = st_mod.extract_tables("xlsx", _workbook([wide, row]))
    assert len(table[0]) == st_mod.MAX_XLSX_COLUMNS
    assert table[1] == row[:5]  # the empty tail of a row is not kept


def test_xlsx_empty_rows_far_apart_cost_no_cells() -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.append(VTB_HEADER)
    sheet.cell(row=3000, column=1, value="03.10.2026")  # thousands of empty rows in between
    buffer = io.BytesIO()
    book.save(buffer)
    (table,) = st_mod.extract_tables("xlsx", buffer.getvalue())
    assert len(table) == 3000
    assert sum(len(row) for row in table) == len(VTB_HEADER) + 1


def test_xlsx_too_many_cells(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(st_mod, "MAX_CELLS", 20)
    rows = [VTB_HEADER] + [["03.10.2026", "x", "1", "", "RUB", "", ""]] * 5
    with pytest.raises(StatementError) as raised:
        parse_statement(_workbook(rows))
    assert raised.value.code == "statement_too_large"
    assert "cells" in raised.value.message


def test_xlsx_too_many_sheets(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(st_mod, "MAX_XLSX_SHEETS", 1)
    with pytest.raises(StatementError) as raised:
        parse_statement(_workbook([VTB_HEADER], extra_sheet=True))
    assert (raised.value.code, "sheets" in raised.value.message) == ("statement_too_large", True)


def test_a_parse_past_its_deadline_stops(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(st_mod, "PARSE_SECONDS", -1.0)
    for data in (_workbook([VTB_HEADER]), PDF, CSV):
        with pytest.raises(StatementError) as raised:
            parse_statement(data)
        assert raised.value.code == "statement_too_large"
        assert "too long" in raised.value.message


def test_too_many_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(st_mod, "MAX_ROWS", 3)
    with pytest.raises(StatementError) as raised:
        parse_statement(CSV)
    assert raised.value.code == "statement_too_large"
    with pytest.raises(StatementError) as xlsx:
        parse_statement(_workbook([VTB_HEADER] + [["03.10.2026", "x", "1", "", "RUB", "", ""]] * 5))
    assert xlsx.value.code == "statement_too_large"
    with pytest.raises(StatementError) as pdf:
        parse_statement(PDF)
    assert pdf.value.code == "statement_too_large"


def test_too_many_pdf_pages(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(st_mod, "MAX_PDF_PAGES", 0)
    with pytest.raises(StatementError) as raised:
        parse_statement(PDF)
    assert raised.value.code == "statement_too_large"


# ------------------------------------------------------------------ formats and errors


def test_format_detection() -> None:
    assert st_mod.detect_format(PDF) == "pdf"
    assert st_mod.detect_format(b"PK\x03\x04...") == "xlsx"
    assert st_mod.detect_format(CSV) == "csv"
    assert st_mod.detect_format(CSV, "csv") == "csv"
    with pytest.raises(StatementError) as raised:
        st_mod.detect_format(CSV, "pdf")
    assert raised.value.code == "statement_format_mismatch"
    with pytest.raises(StatementError) as unknown:
        st_mod.extract_tables("doc", b"")
    assert unknown.value.code == "statement_unsupported"


def test_binary_and_headerless_csv_are_rejected() -> None:
    with pytest.raises(StatementError) as binary:
        parse_statement(b"\x00\x01\x02" * 100)
    assert binary.value.code == "statement_unreadable"
    with pytest.raises(StatementError) as undecodable:
        parse_statement(b"\xff\xfe\x98\x98 \x98")  # neither utf-8 nor cp1251
    assert undecodable.value.code in {"statement_unreadable", "statement_unrecognized"}
    with pytest.raises(StatementError) as none:
        parse_statement(b"hello;world\n1;2\n")
    assert none.value.code == "statement_unrecognized"


# ------------------------------------------------------------------ row rules


def _csv(*rows: str, header: str = "Дата;Сумма;Описание;Валюта") -> bytes:
    return "\n".join([header, *rows]).encode()


def test_row_rules() -> None:
    result = parse_statement(
        _csv(
            "03.10.2026 10:00;-100,00;Кофе;RUB",
            "не дата;-5;Что-то;RUB",
            "не дата;;Подвал;",
            "31.02.2026;-5;Нет такого дня;RUB",
            "01.01.2010;-5;Слишком давно;RUB",
            "03.10.26;+1 000,50 ₽;Двузначный год;RUB",
            "2026-10-04 09:00:00;(75,00);Скобки;RUB",
            "04.10.2026;0;Ноль;RUB",
            "05.10.2026;abc;Не число;RUB",
            "06.10.2026;-1000000000000000;Огромная;RUB",
            "07.10.2026;-9,99;Евро;EUR",
            "07.10.2026;-1,00;Рубли;руб",
            "07.10.2026;-1,00;Рубли код;643",
        )
    )
    reasons = {(s["row"], s["reason"]) for s in result["skipped"]}
    assert (3, "bad_date") in reasons  # has an amount: reported
    assert not any(r == 4 for r, _ in reasons)  # an empty-amount footer is silent
    assert {(5, "bad_date"), (6, "bad_date")} <= reasons
    assert {(9, "no_amount"), (10, "no_amount")} <= reasons
    rows = {c["merchant"]: c for c in candidates(result)}
    assert rows["Кофе"]["amount"] == 10_000
    assert (rows["Двузначный год"]["kind"], rows["Двузначный год"]["amount"]) == ("income", 100_050)
    assert rows["Двузначный год"]["occurred_at"] == "2026-10-03T09:00:00Z"
    assert (rows["Скобки"]["kind"], rows["Скобки"]["amount"]) == ("expense", 7_500)
    assert (rows["Евро"]["currency"], rows["Евро"]["needs_review"]) == ("EUR", True)
    assert (rows["Рубли"]["currency"], rows["Рубли код"]["currency"]) == ("RUB", "RUB")


def test_repeated_header_rows_between_pdf_pages_are_ignored() -> None:
    header = "Дата;Сумма;Описание"
    result = parse_statement(_csv("03.10.2026;-1;A", header, "04.10.2026;-2;B", header=header))
    assert [c["merchant"] for c in candidates(result)] == ["A", "B"]


def test_a_long_merchant_is_cut_and_the_card_tail_is_found() -> None:
    long_name = "Я" * 300
    result = parse_statement(
        _csv(
            f"03.10.2026;-1;{long_name};RUB;**** **** **** 4321",
            "03.10.2026;-1;Short;RUB;нет",
            header="Дата;Сумма;Описание;Валюта;Карта",
        )
    )
    first, second = candidates(result)
    assert len(first["merchant"]) == 200
    assert (first["card_last4"], second["card_last4"]) == ("4321", None)


def test_user_rules_feed_the_suggestions() -> None:
    rules = [
        {
            "id": "r1",
            "merchant_key": "кофе",
            "match_type": "exact",
            "kind": "expense",
            "category_id": "00000000-0000-7000-8000-00000000c001",
        }
    ]
    result = parse_statement(_csv("03.10.2026;-100;Кофе;RUB"), user_rules=rules)
    assert candidates(result)[0]["suggested_category"]["source"] == "user"


def test_closing_balance_needs_a_day_to_hang_on() -> None:
    rows = [["Остаток на конец периода: 100 RUB"], ["Дата", "Сумма"]]
    parsed = st_mod.parse_table(rows)
    assert parsed is not None
    assert parsed["closing_balance"] is None
    assert st_mod.parse_table([["a"], ["b"]]) is None


def test_unreadable_closing_balance_is_ignored() -> None:
    rows = [["Остаток на конец периода: 1.2.3"], ["Дата", "Сумма"], ["03.10.2026", "-5"]]
    parsed = st_mod.parse_table(rows)
    assert parsed is not None
    assert parsed["closing_balance"] is None


def test_cells() -> None:
    assert st_mod.parse_money_cell("") is None
    assert st_mod.parse_money_cell("1 234,56") == 123_456
    assert st_mod.parse_money_cell("−5") == -500
    assert st_mod.parse_money_cell("(1,5)") == -150
    assert st_mod.parse_money_cell("x") is None
    assert st_mod.parse_statement_datetime("03.10.2026 24:00") is None
    assert st_mod.parse_statement_datetime("2026-10-03") == ("2026-10-03", None)
    assert st_mod.parse_statement_datetime("2026-10-03T10:20") == ("2026-10-03", "10:20:00")
    assert st_mod.normalize_currency("") is None
    assert st_mod.normalize_currency("руб.") == "RUB"
    assert st_mod.normalize_currency("chfxxxxxxxxxxxx") == "CHFXXXXX"
    assert st_mod._clean_cell(True) == ""
    assert st_mod._clean_cell(0.1) == "0.1"
    assert st_mod._clean_cell(5) == "5"
    assert st_mod.header_key("  Сумма в валюте СЧЁТА* ") == "сумма в валюте счета"


# ------------------------------------------------------------------ properties


@given(st.text(max_size=60))
def test_normalize_merchant_is_idempotent(text: str) -> None:
    once = ref.normalize_merchant(text)
    assert ref.normalize_merchant(once) == once
    assert once == once.strip()
    assert "  " not in once


@given(st.text(max_size=40), st.text(max_size=40))
def test_similarity_is_symmetric_and_bounded(a: str, b: str) -> None:
    na, nb = ref.normalize_merchant(a), ref.normalize_merchant(b)
    value = ref.similarity(na, nb)
    assert 0 <= value <= 100
    assert value == ref.similarity(nb, na)
    if na:
        assert ref.similarity(na, na) == 100
